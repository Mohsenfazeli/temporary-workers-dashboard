"""Shared python-docx helpers for the Downloads note Word document."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx.opc.constants as C

NAVY = RGBColor(0x00, 0x33, 0x66)
GOLD_HEX = "FCBA19"
GREY = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"


def set_font(style, size=None, bold=None, color=None, italic=None, name=FONT):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size is not None: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color


def para_fmt(style, before=0, after=3, line=1.0, keep_next=None):
    pf = style.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if keep_next is not None: pf.keep_with_next = keep_next


def new_document(margins=(0.55, 0.6, 0.55, 0.6), body_pt=10.5):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = [Inches(m) for m in margins]
    sec.footer_distance = Inches(0.3); sec.header_distance = Inches(0.3)
    st = doc.styles
    set_font(st["Normal"], body_pt); para_fmt(st["Normal"], 0, 3, 1.0)
    set_font(st["Title"], 16, True, NAVY); para_fmt(st["Title"], 0, 1, 1.0)
    # remove the default Title bottom border/spacing quirks
    tppr = st["Title"].element.get_or_add_pPr()
    for b in tppr.findall(qn("w:pBdr")): tppr.remove(b)
    set_font(st["Heading 1"], 12, True, NAVY); para_fmt(st["Heading 1"], 6, 1.5, 1.0, True)
    set_font(st["Heading 2"], 10.5, True, NAVY); para_fmt(st["Heading 2"], 4, 1, 1.0, True)
    set_font(st["Heading 3"], 10, True, NAVY); para_fmt(st["Heading 3"], 3, 1, 1.0, True)
    set_font(st["Caption"], 9.5, True, NAVY, italic=False); para_fmt(st["Caption"], 4, 1, 1.0, True)
    set_font(st["List Bullet"], body_pt); para_fmt(st["List Bullet"], 0, 1.5, 1.0)
    for name, size, color in (("Source Note", 8, GREY), ("Byline", 9, GREY), ("Reference", 8, None),
                              ("Table Text", 8, None), ("Small Note", 8.5, None)):
        s = st.add_style(name, 1); s.base_style = st["Normal"]
        set_font(s, size, color=color); para_fmt(s, 0, 1 if name != "Byline" else 4, 1.0)
    return doc


def add_field(paragraph, instr, display, style_bold=None):
    r1 = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r1._r.append(fc)
    r2 = paragraph.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r2._r.append(it)
    r3 = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r3._r.append(fc)
    r4 = paragraph.add_run(display)
    if style_bold is not None: r4.bold = style_bold
    r5 = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r5._r.append(fc)
    return r4


def add_hyperlink(paragraph, url, text, size=None, color="0B3D6B"):
    part = paragraph.part
    r_id = part.relate_to(url, C.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); paragraph._p.append(h)


def shade(paragraph, fill="EEF3F8", left_color="003366", left_sz=18, box=False):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    bdr = OxmlElement("w:pBdr")
    sides = ("top", "left", "bottom", "right") if box else ("left",)
    for side in sides:
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(left_sz if side == "left" else 4)); e.set(qn("w:space"), "4")
        e.set(qn("w:color"), left_color if side == "left" else "C9D6E3"); bdr.append(e)
    # pBdr must come before shd in pPr sequence
    ppr.append(bdr); ppr.append(shd)


def rule(paragraph, color=GOLD_HEX, sz=12):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "1"); b.set(qn("w:color"), color)
    bdr.append(b); ppr.append(bdr)


TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def rich(paragraph, text, size=None):
    """Add text with **bold** and *italic* markup."""
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and len(part) > 1:
            r = paragraph.add_run(part[1:-1]); r.italic = True
        else:
            r = paragraph.add_run(part)
        if size: r.font.size = Pt(size)
    return paragraph


def footer(doc, text):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]; p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(text + " · page "); r.font.size = Pt(8); r.font.color.rgb = GREY
        x = add_field(p, "PAGE", "1"); x.font.size = Pt(8); x.font.color.rgb = GREY


def caption(doc, label, title, number):
    p = doc.add_paragraph(style="Caption")
    p.add_run(f"{label} ")
    add_field(p, f"SEQ {label.replace(' ', '_')} \\* ARABIC", str(number))
    p.add_run(f". {title}")
    return p


def table(doc, header, rows, widths, font_pt=8, header_fill="003366", header_text="FFFFFF", stripe="F3F6F9"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = doc.styles["Table Grid"]; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        p = c.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.size = Pt(font_pt); r.font.color.rgb = RGBColor.from_string(header_text)
        tcpr = c._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), header_fill); tcpr.append(shd)
    # repeat header row
    trpr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            rich(p, str(v), size=font_pt)
            if stripe and k % 2 == 1:
                tcpr = cells[i]._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), stripe); tcpr.append(shd)
    tblpr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    return t


BLACK = RGBColor(0, 0, 0)
DGREY = RGBColor(0x40, 0x40, 0x40)


def new_plain_document(margins=(0.75, 0.8, 0.75, 0.8), body_pt=11):
    """Plain government-report look: Calibri, black text, bold black headings,
    no rules, shading or colour."""
    doc = new_document(margins=margins, body_pt=body_pt)
    st = doc.styles
    set_font(st["Normal"], body_pt, color=BLACK); para_fmt(st["Normal"], 0, 5, 1.05)
    set_font(st["Title"], 15, True, BLACK); para_fmt(st["Title"], 0, 3, 1.0)
    set_font(st["Heading 1"], 12.5, True, BLACK); para_fmt(st["Heading 1"], 8, 3, 1.0, True)
    set_font(st["Heading 2"], 11, True, BLACK); para_fmt(st["Heading 2"], 6, 2, 1.0, True)
    set_font(st["Heading 3"], 11, True, BLACK, italic=False); para_fmt(st["Heading 3"], 4, 2, 1.0, True)
    set_font(st["Caption"], 10, True, BLACK, italic=False); para_fmt(st["Caption"], 6, 2, 1.0, True)
    set_font(st["List Bullet"], body_pt, color=BLACK); para_fmt(st["List Bullet"], 0, 3, 1.05)
    for name, size in (("Source Note", 9), ("Byline", 10), ("Reference", 9.5), ("Table Text", 9), ("Small Note", 10)):
        set_font(st[name], size, color=DGREY if name in ("Source Note", "Byline") else BLACK)
    # strip theme colours from the heading/title styles so they are plain black
    for sname in ("Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        rpr = st[sname].element.get_or_add_rPr()
        c = rpr.find(qn("w:color"))
        if c is not None:
            for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
                if c.get(qn(a)) is not None: del c.attrib[qn(a)]
    return doc


def plain_footer(doc, text):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]; p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(text + " | page "); r.font.size = Pt(9); r.font.color.rgb = DGREY
        x = add_field(p, "PAGE", "1"); x.font.size = Pt(9); x.font.color.rgb = DGREY


def plain_table(doc, header, rows, widths, font_pt=9):
    return table(doc, header, rows, widths, font_pt=font_pt, header_fill="F2F2F2", header_text="000000", stripe=None)
